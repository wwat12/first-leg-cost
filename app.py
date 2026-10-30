#!/usr/bin/env python3
"""Amazon FBA 头程运费分摊 MVP — Flask app."""
from __future__ import annotations

import json
import os
import uuid
from datetime import datetime
from pathlib import Path

from flask import (
    Flask,
    flash,
    redirect,
    render_template,
    request,
    send_file,
    session,
    url_for,
    jsonify,
)
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

from lib.bill_parser import parse_bill, detect_detail_sheets
from lib.saihu import peek_columns, parse_saihu
from lib.allocate import allocate_by_weight, freight_by_shipment_export_rows

BASE = Path(__file__).resolve().parent
UPLOADS = BASE / "uploads"
EXPORTS = BASE / "exports"
UPLOADS.mkdir(exist_ok=True)
EXPORTS.mkdir(exist_ok=True)

app = Flask(__name__)
app.secret_key = os.environ.get("FLASK_SECRET", "first-leg-cost-dev-key")
app.config["MAX_CONTENT_LENGTH"] = 64 * 1024 * 1024  # 64MB


def _sid() -> str:
    if "sid" not in session:
        session["sid"] = uuid.uuid4().hex
    return session["sid"]


def _state_path() -> Path:
    return UPLOADS / f"{_sid()}_state.json"


def _load_state() -> dict:
    p = _state_path()
    if p.exists():
        return json.loads(p.read_text(encoding="utf-8"))
    return {}


def _save_state(state: dict) -> None:
    _state_path().write_text(json.dumps(state, ensure_ascii=False, default=str), encoding="utf-8")


def _thin_border():
    s = Side(style="thin", color="CCCCCC")
    return Border(left=s, right=s, top=s, bottom=s)


def _write_sheet(ws, rows: list[dict], headers: list[str] | None = None):
    if not rows:
        ws.append(["（无数据）"])
        return
    headers = headers or list(rows[0].keys())
    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(color="FFFFFF", bold=True)
    for col, h in enumerate(headers, 1):
        cell = ws.cell(1, col, h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")
        cell.border = _thin_border()
    for r_i, row in enumerate(rows, 2):
        for c_i, h in enumerate(headers, 1):
            val = row.get(h)
            if isinstance(val, (list, set)):
                val = ",".join(str(x) for x in val)
            cell = ws.cell(r_i, c_i, val)
            cell.border = _thin_border()
    for c in range(1, len(headers) + 1):
        ws.column_dimensions[ws.cell(1, c).column_letter].width = 16


@app.route("/")
def index():
    state = _load_state()
    bill = state.get("bill")
    saihu = state.get("saihu")
    alloc = state.get("alloc")
    return render_template(
        "index.html",
        shop=state.get("shop", ""),
        bill=bill,
        saihu=saihu,
        alloc=alloc,
        saihu_peek=state.get("saihu_peek"),
        saihu_mapping=state.get("saihu_mapping"),
    )


@app.route("/upload_bill", methods=["POST"])
def upload_bill():
    shop = (request.form.get("shop") or "").strip()
    f = request.files.get("bill_file")
    if not f or not f.filename:
        flash("请选择物流账单 Excel 文件", "error")
        return redirect(url_for("index"))
    sid = _sid()
    dest = UPLOADS / f"{sid}_bill.xlsx"
    f.save(dest)
    try:
        result = parse_bill(str(dest), shop=shop)
    except Exception as e:
        flash(f"解析账单失败: {e}", "error")
        return redirect(url_for("index"))
    if not result.get("rows"):
        flash("未解析到运费行，请确认文件含明细表且费用类型含「运费」", "error")
    else:
        flash(
            f"账单解析成功：运费合计 ¥{result['total_freight']:.2f}，"
            f"{result['freight_row_count']} 行 / {result['ticket_count']} 票",
            "ok",
        )
    state = _load_state()
    state["shop"] = shop
    state["bill"] = result
    state["bill_path"] = str(dest)
    # clear downstream
    state.pop("alloc", None)
    _save_state(state)
    return redirect(url_for("index"))


@app.route("/upload_saihu", methods=["POST"])
def upload_saihu():
    f = request.files.get("saihu_file")
    if not f or not f.filename:
        flash("请选择赛狐货件导出 Excel", "error")
        return redirect(url_for("index"))
    sid = _sid()
    dest = UPLOADS / f"{sid}_saihu.xlsx"
    f.save(dest)
    try:
        peek = peek_columns(str(dest))
    except Exception as e:
        flash(f"读取赛狐文件失败: {e}", "error")
        return redirect(url_for("index"))
    state = _load_state()
    state["saihu_path"] = str(dest)
    state["saihu_peek"] = peek
    state.pop("saihu", None)
    state.pop("alloc", None)
    # prefill suggested mapping from first sheet
    if peek.get("sheets"):
        state["saihu_mapping"] = peek["sheets"][0].get("suggested_mapping") or {}
        state["saihu_sheet"] = peek["sheets"][0]["name"]
        state["saihu_header_row"] = peek["sheets"][0]["header_row"]
    _save_state(state)
    flash("赛狐文件已上传，请确认列映射后点击「解析赛狐数据」", "ok")
    return redirect(url_for("index") + "#saihu")


@app.route("/parse_saihu", methods=["POST"])
def parse_saihu_route():
    state = _load_state()
    path = state.get("saihu_path")
    if not path or not Path(path).exists():
        flash("请先上传赛狐文件", "error")
        return redirect(url_for("index"))
    sheet = request.form.get("sheet") or state.get("saihu_sheet")
    header_row = int(request.form.get("header_row") or state.get("saihu_header_row") or 1)
    weight_mode = request.form.get("weight_mode") or "unit"
    mapping = {
        "fba_id": request.form.get("map_fba_id") or "",
        "sku": request.form.get("map_sku") or "",
        "qty": request.form.get("map_qty") or "",
        "unit_weight": request.form.get("map_unit_weight") or "",
        "total_weight": request.form.get("map_total_weight") or "",
    }
    result = parse_saihu(
        path,
        mapping=mapping,
        sheet_name=sheet,
        header_row=header_row,
        weight_mode=weight_mode,
    )
    if not result.get("ok"):
        flash(result.get("error") or "解析失败", "error")
        state["saihu_mapping"] = mapping
        _save_state(state)
        return redirect(url_for("index") + "#saihu")
    state["saihu"] = result
    state["saihu_mapping"] = mapping
    state["saihu_sheet"] = sheet
    state["saihu_header_row"] = header_row
    state["saihu_weight_mode"] = weight_mode
    state.pop("alloc", None)
    _save_state(state)
    flash(
        f"赛狐解析成功：{result['row_count']} 行，"
        f"{result['fba_count']} 个货件，{result['sku_count']} 个 SKU",
        "ok",
    )
    return redirect(url_for("index") + "#allocate")


@app.route("/allocate", methods=["POST"])
def allocate_route():
    state = _load_state()
    bill = state.get("bill")
    saihu = state.get("saihu")
    if not bill or not bill.get("rows"):
        flash("请先上传并解析物流账单", "error")
        return redirect(url_for("index"))
    if not saihu or not saihu.get("rows"):
        flash("请先上传并解析赛狐货件数据", "error")
        return redirect(url_for("index") + "#saihu")
    shop = state.get("shop") or ""
    result = allocate_by_weight(bill["rows"], saihu["rows"], shop=shop)
    state["alloc"] = result
    _save_state(state)
    flash(
        f"分摊完成：{result['sku_count']} 个 SKU，"
        f"已分摊运费 ¥{result['total_allocated']:.2f}，"
        f"匹配货件 {result['matched_fba_count']} 个",
        "ok",
    )
    return redirect(url_for("index") + "#result")


@app.route("/export/freight")
def export_freight():
    state = _load_state()
    bill = state.get("bill")
    if not bill or not bill.get("rows"):
        flash("无账单数据可导出", "error")
        return redirect(url_for("index"))
    shop = state.get("shop") or ""
    rows = freight_by_shipment_export_rows(bill["rows"], shop=shop)
    wb = Workbook()
    ws = wb.active
    ws.title = "运费按货件"
    _write_sheet(ws, rows)
    # tickets sheet
    ws2 = wb.create_sheet("按运单汇总")
    ticket_rows = []
    for t in bill.get("tickets") or []:
        ticket_rows.append(
            {
                "店铺": shop,
                "运单号": t["waybill_no"],
                "运费合计": round(t["freight_total"], 4),
                "收费重合计": round(t["weight_total"], 4),
                "运费行数": t["line_count"],
                "FBA货件号": ",".join(t.get("fba_ids") or []),
                "客户单号": ",".join(t.get("customer_nos") or []),
            }
        )
    _write_sheet(ws2, ticket_rows)
    # summary
    ws3 = wb.create_sheet("汇总")
    _write_sheet(
        ws3,
        [
            {
                "店铺": shop,
                "明细表": bill.get("sheet"),
                "运费行数": bill.get("freight_row_count"),
                "运单数": bill.get("ticket_count"),
                "运费合计_CNY": bill.get("total_freight"),
                "收费重合计": bill.get("total_weight"),
                "排除费用类型": json.dumps(bill.get("skipped_fee_types") or {}, ensure_ascii=False),
                "导出时间": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            }
        ],
    )
    out = EXPORTS / f"freight_{_sid()[:8]}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    wb.save(out)
    return send_file(out, as_attachment=True, download_name=out.name)


@app.route("/export/sku")
def export_sku():
    state = _load_state()
    alloc = state.get("alloc")
    if not alloc or not alloc.get("sku_results"):
        flash("请先完成分摊", "error")
        return redirect(url_for("index"))
    shop = state.get("shop") or ""
    wb = Workbook()
    ws = wb.active
    ws.title = "SKU单件头程"
    sku_rows = []
    for r in alloc["sku_results"]:
        sku_rows.append(
            {
                "店铺": shop or r.get("shop") or "",
                "SKU": r["sku"],
                "总发货数量": r["total_qty"],
                "分摊运费合计": r["total_freight"],
                "单件头程": r["unit_first_leg"],
                "涉及货件数": r["shipment_count"],
                "FBA货件号": ",".join(r.get("fba_ids") or []),
            }
        )
    _write_sheet(ws, sku_rows)

    ws2 = wb.create_sheet("货件SKU明细")
    detail_rows = []
    for r in alloc.get("shipment_allocs") or []:
        detail_rows.append(
            {
                "店铺": shop,
                "运单号": r["waybill_no"],
                "FBA货件号": r["fba_id"],
                "SKU": r["sku"],
                "数量": r["qty"],
                "分摊重量": r["alloc_weight"],
                "货件运费": r["shipment_freight"],
                "SKU分摊费用": r["sku_freight"],
                "本行单件头程": r["unit_first_leg_partial"],
            }
        )
    _write_sheet(ws2, detail_rows)

    if alloc.get("unmatched_freight"):
        ws3 = wb.create_sheet("未匹配运费")
        _write_sheet(
            ws3,
            [
                {
                    "运单号": u["waybill_no"],
                    "运费": u["freight"],
                    "FBA货件号": ",".join(u.get("fba_ids") or []),
                    "原因": u.get("reason"),
                }
                for u in alloc["unmatched_freight"]
            ],
        )

    out = EXPORTS / f"sku_firstleg_{_sid()[:8]}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    wb.save(out)
    return send_file(out, as_attachment=True, download_name=out.name)


@app.route("/api/bill_summary")
def api_bill_summary():
    state = _load_state()
    bill = state.get("bill")
    if not bill:
        return jsonify({"ok": False, "error": "no bill"})
    return jsonify(
        {
            "ok": True,
            "total_freight": bill.get("total_freight"),
            "ticket_count": bill.get("ticket_count"),
            "freight_row_count": bill.get("freight_row_count"),
            "sheet": bill.get("sheet"),
        }
    )


@app.route("/reset", methods=["POST"])
def reset():
    sid = _sid()
    for p in UPLOADS.glob(f"{sid}*"):
        try:
            p.unlink()
        except OSError:
            pass
    session.clear()
    flash("已清空会话数据", "ok")
    return redirect(url_for("index"))


if __name__ == "__main__":
    port = int(os.environ.get("PORT", "8765"))
    print(f"头程运费分摊工具: http://127.0.0.1:{port}")
    app.run(host="0.0.0.0", port=port, debug=os.environ.get("FLASK_DEBUG", "0") == "1")
