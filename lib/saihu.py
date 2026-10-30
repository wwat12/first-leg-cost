"""Parse Saihu (赛狐) FBA exports — prefer 装箱信息 packing sheet."""
from __future__ import annotations

from collections import defaultdict
from typing import Any

from openpyxl import load_workbook


FBA_ALIASES = [
    "货件编号", "货件号", "FBA货件号", "FBA单号", "Shipment ID", "ShipmentId",
    "shipment_id", "亚马逊货件号", "FBA shipment", "ShipmentID",
]
SKU_ALIASES = [
    "SKU", "卖家SKU", "MSKU", "Seller SKU", "商家SKU", "sku", "产品SKU",
]
QTY_ALIASES = [
    "数量", "发货数量", "Quantity", "Qty", "件数", "发货件数", "计划发货量", "实际发货量",
]
QTY_PER_BOX_ALIASES = ["单箱数量", "每箱数量", "箱内数量"]
BOX_COUNT_ALIASES = ["箱数", "箱子数"]
UNIT_WEIGHT_ALIASES = [
    "单件重量", "单件含包装重量", "单重", "单位重量", "Unit Weight", "unit_weight",
    "产品重量", "毛重(单件)", "单件毛重",
]
TOTAL_WEIGHT_ALIASES = [
    "总重量", "箱子重量(kg)", "箱子重量", "重量", "Weight", "合计重量", "发货重量", "毛重",
]
SHOP_ALIASES = ["店铺", "店铺名称", "Shop", "store", "站点店铺"]


def _norm(v: Any) -> str:
    if v is None:
        return ""
    return str(v).strip()


def _to_float(v: Any) -> float:
    if v is None or v == "":
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(",", "").replace("，", "")
    try:
        return float(s)
    except ValueError:
        return 0.0


def _find_col(headers: dict[str, int], aliases: list[str]) -> str | None:
    for a in aliases:
        if a in headers:
            return a
    lower = {k.lower(): k for k in headers}
    for a in aliases:
        if a.lower() in lower:
            return lower[a.lower()]
    for k in headers:
        for a in aliases:
            if a.lower() in k.lower():
                return k
    return None


def _header_map(ws, header_row: int) -> dict[str, int]:
    headers: dict[str, int] = {}
    max_col = ws.max_column or 1
    for c in range(1, max_col + 1):
        val = _norm(ws.cell(header_row, c).value)
        if val:
            headers[val] = c
    return headers


def _detect_header_row(ws, max_scan: int = 15) -> tuple[int, dict[str, int]]:
    best_i, best_score, best_headers = 1, -1, {}
    max_row = min(ws.max_row or 1, max_scan)
    for r in range(1, max_row + 1):
        headers = _header_map(ws, r)
        if not headers:
            continue
        score = len(headers)
        keys = set(headers)
        for alias_list in (FBA_ALIASES, SKU_ALIASES, TOTAL_WEIGHT_ALIASES, QTY_PER_BOX_ALIASES):
            if any(a in keys for a in alias_list):
                score += 5
        if "货件编号" in keys:
            score += 10
        if score > best_score:
            best_i, best_score, best_headers = r, score, headers
    return best_i, best_headers


def peek_columns(path: str, max_header_scan: int = 15) -> dict[str, Any]:
    """Return sheet names, detected header row, columns, samples, suggested mapping."""
    wb = load_workbook(path, data_only=True, read_only=True)
    try:
        sheets_info = []
        for name in wb.sheetnames:
            ws = wb[name]
            rows = []
            for i, row in enumerate(ws.iter_rows(max_row=30, max_col=40, values_only=True)):
                rows.append([_norm(v) for v in row])
                if i >= 29:
                    break
            best = (-1, 0, [])
            for i, row in enumerate(rows[:max_header_scan]):
                nonempty = [c for c in row if c]
                score = len(nonempty)
                joined = set(nonempty)
                for alias_list in (FBA_ALIASES, SKU_ALIASES, QTY_ALIASES, QTY_PER_BOX_ALIASES):
                    if any(a in joined for a in alias_list):
                        score += 5
                if "货件编号" in joined:
                    score += 10
                if "装箱" in name:
                    score += 3
                if score > best[1]:
                    best = (i, score, nonempty)
            header_row_idx, _, cols = best
            sample = []
            for row in rows[header_row_idx + 1 : header_row_idx + 6]:
                if any(row):
                    sample.append(row[: len(cols)] if cols else row)
            suggested = suggest_mapping(cols)
            sheets_info.append(
                {
                    "name": name,
                    "header_row": header_row_idx + 1,
                    "columns": cols,
                    "sample": sample,
                    "suggested_mapping": suggested,
                    "is_packing": "装箱" in name,
                }
            )
        # Prefer 装箱信息 sheet first in UI
        sheets_info.sort(key=lambda s: (0 if s.get("is_packing") else 1, s["name"]))
        return {"ok": True, "sheets": sheets_info}
    finally:
        wb.close()


def suggest_mapping(columns: list[str]) -> dict[str, str | None]:
    def find(aliases: list[str], exclude: set[str] | None = None) -> str | None:
        exclude = exclude or set()
        candidates = [c for c in columns if c not in exclude]
        for c in candidates:
            if c in aliases:
                return c
        lower_map = {c.lower(): c for c in candidates}
        for a in aliases:
            if a.lower() in lower_map:
                return lower_map[a.lower()]
        for c in candidates:
            for a in aliases:
                if a.lower() in c.lower():
                    return c
        return None

    unit_w = find(UNIT_WEIGHT_ALIASES)
    total_w = find(TOTAL_WEIGHT_ALIASES, exclude={unit_w} if unit_w else set())
    if total_w and unit_w and total_w == unit_w:
        total_w = None
    # Prefer SKU over MSKU when both exist
    sku = None
    if "SKU" in columns:
        sku = "SKU"
    else:
        sku = find(SKU_ALIASES)
    return {
        "fba_id": find(FBA_ALIASES),
        "sku": sku,
        "qty": find(QTY_ALIASES),
        "qty_per_box": find(QTY_PER_BOX_ALIASES),
        "box_count": find(BOX_COUNT_ALIASES),
        "unit_weight": unit_w,
        "total_weight": total_w,
        "shop": find(SHOP_ALIASES),
    }


def _load_shop_by_fba(wb) -> dict[str, str]:
    """From 货件信息 (or any sheet with 货件编号+店铺), map FBA -> shop."""
    out: dict[str, str] = {}
    for name in wb.sheetnames:
        ws = wb[name]
        header_row, headers = _detect_header_row(ws)
        fba_col = _find_col(headers, FBA_ALIASES)
        shop_col = _find_col(headers, SHOP_ALIASES)
        if not fba_col or not shop_col:
            continue
        for r in range(header_row + 1, (ws.max_row or header_row) + 1):
            fba = _norm(ws.cell(r, headers[fba_col]).value)
            shop = _norm(ws.cell(r, headers[shop_col]).value)
            if fba and shop and fba not in out:
                out[fba] = shop
    return out


def parse_packing(
    path: str,
    sheet_name: str | None = None,
    header_row: int | None = None,
    unit_weights: dict[str, float] | None = None,
) -> dict[str, Any]:
    """
    Parse 装箱信息-style sheet.

    weight (货件实际毛重) = 总重量, else 箱子重量(kg) aggregated per FBA.
    qty = 单箱数量 × 箱数 (箱数 missing → 单箱数量).
    Shop joined from 货件信息 when available.
    """
    unit_weights = unit_weights or {}
    wb = load_workbook(path, data_only=True)
    try:
        # Prefer 装箱信息
        if sheet_name and sheet_name in wb.sheetnames:
            name = sheet_name
        else:
            name = None
            for n in wb.sheetnames:
                if "装箱" in n:
                    name = n
                    break
            if name is None:
                name = wb.sheetnames[0]
        ws = wb[name]
        if header_row is None:
            header_row, headers = _detect_header_row(ws)
        else:
            headers = _header_map(ws, header_row)

        fba_col = _find_col(headers, FBA_ALIASES)
        sku_col = _find_col(headers, ["SKU"]) or _find_col(headers, SKU_ALIASES)
        qty_per_box_col = _find_col(headers, QTY_PER_BOX_ALIASES)
        box_count_col = _find_col(headers, BOX_COUNT_ALIASES)
        qty_col = _find_col(headers, QTY_ALIASES)
        total_w_col = _find_col(headers, ["总重量"]) or _find_col(headers, TOTAL_WEIGHT_ALIASES)
        box_w_col = _find_col(headers, ["箱子重量(kg)", "箱子重量"])
        unit_w_col = _find_col(headers, UNIT_WEIGHT_ALIASES)
        shop_col = _find_col(headers, SHOP_ALIASES)

        if not fba_col or not sku_col:
            return {
                "ok": False,
                "error": f"装箱表缺少货件编号或 SKU 列（表：{name}）",
                "available_columns": list(headers.keys()),
            }
        if not qty_per_box_col and not qty_col:
            return {
                "ok": False,
                "error": "装箱表缺少数量列（需要「单箱数量」或「数量」）",
                "available_columns": list(headers.keys()),
            }

        shop_by_fba = _load_shop_by_fba(wb)

        # Accumulate per (fba, sku)
        # shipment_weight: prefer 总重量 (once per fba), else sum box weights
        ship_total_w: dict[str, float] = {}
        ship_box_w_sum: dict[str, float] = defaultdict(float)
        lines: dict[tuple[str, str], dict[str, Any]] = {}

        for r in range(header_row + 1, (ws.max_row or header_row) + 1):
            fba = _norm(ws.cell(r, headers[fba_col]).value)
            sku = _norm(ws.cell(r, headers[sku_col]).value)
            if not fba and not sku:
                continue
            if not fba or not sku:
                continue

            if qty_per_box_col:
                per_box = _to_float(ws.cell(r, headers[qty_per_box_col]).value)
                boxes = _to_float(ws.cell(r, headers[box_count_col]).value) if box_count_col else 0.0
                if boxes > 0:
                    qty = per_box * boxes
                else:
                    qty = per_box
            else:
                qty = _to_float(ws.cell(r, headers[qty_col]).value)

            if total_w_col:
                tw = _to_float(ws.cell(r, headers[total_w_col]).value)
                if tw > 0 and fba not in ship_total_w:
                    ship_total_w[fba] = tw
            if box_w_col:
                bw = _to_float(ws.cell(r, headers[box_w_col]).value)
                if bw > 0:
                    # If 箱数>1 and 箱子重量 looks per-box, multiply; sample data has
                    # 箱子重量 == 总重量 when 箱数=1. Prefer adding row weight as-is
                    # when 总重量 missing; when 箱数 present and total missing, * boxes.
                    boxes = _to_float(ws.cell(r, headers[box_count_col]).value) if box_count_col else 1.0
                    if fba not in ship_total_w and boxes > 1 and bw * boxes > bw:
                        ship_box_w_sum[fba] += bw * boxes
                    else:
                        ship_box_w_sum[fba] += bw

            row_unit_w = 0.0
            if unit_w_col:
                row_unit_w = _to_float(ws.cell(r, headers[unit_w_col]).value)
            if row_unit_w <= 0 and sku in unit_weights:
                row_unit_w = float(unit_weights[sku])

            shop = ""
            if shop_col:
                shop = _norm(ws.cell(r, headers[shop_col]).value)
            if not shop:
                shop = shop_by_fba.get(fba, "")

            key = (fba, sku)
            agg = lines.setdefault(
                key,
                {
                    "fba_id": fba,
                    "sku": sku,
                    "qty": 0.0,
                    "unit_weight": 0.0,
                    "shop": shop,
                    "row_count": 0,
                },
            )
            agg["qty"] += qty
            if row_unit_w > 0 and agg["unit_weight"] <= 0:
                agg["unit_weight"] = row_unit_w
            if shop and not agg["shop"]:
                agg["shop"] = shop
            agg["row_count"] += 1

        # Build shipment SKU rows with gross weight per FBA
        skus_by_fba: dict[str, list[dict]] = defaultdict(list)
        for (fba, sku), agg in lines.items():
            gross = ship_total_w.get(fba)
            if not gross or gross <= 0:
                gross = ship_box_w_sum.get(fba, 0.0)
            skus_by_fba[fba].append(
                {
                    "fba_id": fba,
                    "sku": agg["sku"],
                    "qty": agg["qty"],
                    "unit_weight": agg["unit_weight"],
                    "shop": agg["shop"] or "",
                }
            )

        rows: list[dict[str, Any]] = []
        shipment_weights: dict[str, float] = {}
        for fba, sku_list in skus_by_fba.items():
            gross = ship_total_w.get(fba) or ship_box_w_sum.get(fba, 0.0)
            shipment_weights[fba] = float(gross or 0.0)
            for s in sku_list:
                s["shipment_gross_weight"] = shipment_weights[fba]
                rows.append(s)

        shops = sorted({r["shop"] for r in rows if r.get("shop")})
        return {
            "ok": True,
            "sheet": name,
            "header_row": header_row,
            "row_count": len(rows),
            "fba_count": len(skus_by_fba),
            "sku_count": len({r["sku"] for r in rows}),
            "shops": shops,
            "shipment_weights": shipment_weights,
            "rows": rows,
            "source": "packing",
        }
    finally:
        wb.close()


def parse_unit_weight_table(path: str) -> dict[str, Any]:
    """Parse optional SKU → 单件含包装重量 (kg) table."""
    wb = load_workbook(path, data_only=True)
    try:
        ws = wb[wb.sheetnames[0]]
        header_row, headers = _detect_header_row(ws)
        sku_col = _find_col(headers, ["SKU"]) or _find_col(headers, SKU_ALIASES)
        w_col = _find_col(headers, UNIT_WEIGHT_ALIASES) or _find_col(
            headers, ["重量(kg)", "重量", "毛重", "Weight"]
        )
        if not sku_col or not w_col:
            return {
                "ok": False,
                "error": "单件重量表需要 SKU 列与重量列",
                "available_columns": list(headers.keys()),
            }
        weights: dict[str, float] = {}
        for r in range(header_row + 1, (ws.max_row or header_row) + 1):
            sku = _norm(ws.cell(r, headers[sku_col]).value)
            w = _to_float(ws.cell(r, headers[w_col]).value)
            if sku and w > 0:
                weights[sku] = w
        return {"ok": True, "count": len(weights), "weights": weights}
    finally:
        wb.close()


def parse_saihu(
    path: str,
    mapping: dict[str, str] | None = None,
    sheet_name: str | None = None,
    header_row: int = 1,
    weight_mode: str = "unit",
    unit_weights: dict[str, float] | None = None,
) -> dict[str, Any]:
    """
    Main entry: auto-parse 装箱信息 when possible; else use explicit mapping
    (legacy / generic sheet).
    """
    wb = load_workbook(path, data_only=True, read_only=True)
    try:
        names = list(wb.sheetnames)
    finally:
        wb.close()

    # Auto packing path when 装箱 sheet exists and no forced generic mapping
    packing_sheet = sheet_name if sheet_name and "装箱" in sheet_name else None
    if packing_sheet is None:
        for n in names:
            if "装箱" in n:
                packing_sheet = n
                break

    if packing_sheet and (mapping is None or not mapping.get("fba_id")):
        return parse_packing(
            path,
            sheet_name=packing_sheet,
            header_row=None,
            unit_weights=unit_weights,
        )

    # If user pointed at 装箱信息 explicitly even with mapping, prefer packing parser
    if sheet_name and "装箱" in sheet_name:
        return parse_packing(
            path,
            sheet_name=sheet_name,
            header_row=header_row if header_row else None,
            unit_weights=unit_weights,
        )

    # Generic mapped parse (synth / legacy)
    mapping = mapping or {}
    wb = load_workbook(path, data_only=True)
    try:
        name = sheet_name or wb.sheetnames[0]
        ws = wb[name]
        headers = _header_map(ws, header_row)
        required = ["fba_id", "sku", "qty"]
        for k in required:
            col_name = mapping.get(k)
            if not col_name or col_name not in headers:
                # try auto packing as fallback
                if any("装箱" in n for n in wb.sheetnames):
                    return parse_packing(path, unit_weights=unit_weights)
                return {
                    "ok": False,
                    "error": f"缺少列映射: {k} -> {col_name}（表头中未找到）",
                    "available_columns": list(headers.keys()),
                }

        weight_col_name = mapping.get("unit_weight") or mapping.get("total_weight")
        weight_is_total = weight_mode == "total" or (
            mapping.get("total_weight") == weight_col_name and not mapping.get("unit_weight")
        )

        shop_by_fba = _load_shop_by_fba(wb)
        shop_col_name = mapping.get("shop")
        rows = []
        for r in range(header_row + 1, (ws.max_row or header_row) + 1):
            fba = _norm(ws.cell(r, headers[mapping["fba_id"]]).value)
            sku = _norm(ws.cell(r, headers[mapping["sku"]]).value)
            if not fba and not sku:
                continue
            qty_raw = ws.cell(r, headers[mapping["qty"]]).value
            qty = _to_float(qty_raw)
            unit_w = 0.0
            ship_w = 0.0
            if weight_col_name and weight_col_name in headers:
                w = _to_float(ws.cell(r, headers[weight_col_name]).value)
                if weight_is_total:
                    ship_w = w
                    unit_w = (w / qty) if qty else 0.0
                else:
                    unit_w = w
                    ship_w = w * qty
            if unit_w <= 0 and sku in (unit_weights or {}):
                unit_w = float(unit_weights[sku])
                ship_w = unit_w * qty
            shop = ""
            if shop_col_name and shop_col_name in headers:
                shop = _norm(ws.cell(r, headers[shop_col_name]).value)
            if not shop:
                shop = shop_by_fba.get(fba, "")
            rows.append(
                {
                    "fba_id": fba,
                    "sku": sku,
                    "qty": qty,
                    "unit_weight": unit_w,
                    "shipment_gross_weight": ship_w,  # may re-aggregate later
                    "shop": shop,
                    "_weight_is_total_line": weight_is_total,
                }
            )

        # Aggregate shipment gross weight per FBA for generic sheets
        by_fba: dict[str, list] = defaultdict(list)
        for row in rows:
            by_fba[row["fba_id"]].append(row)
        shipment_weights = {}
        out_rows = []
        for fba, lst in by_fba.items():
            if any(x.get("_weight_is_total_line") for x in lst):
                # lines already carry totals — if multi-sku totals are per-SKU line weights
                gross = sum(x["shipment_gross_weight"] for x in lst)
            else:
                gross = sum(x["unit_weight"] * x["qty"] for x in lst)
            shipment_weights[fba] = gross
            for x in lst:
                out_rows.append(
                    {
                        "fba_id": fba,
                        "sku": x["sku"],
                        "qty": x["qty"],
                        "unit_weight": x["unit_weight"],
                        "shipment_gross_weight": gross,
                        "shop": x["shop"],
                    }
                )

        shops = sorted({r["shop"] for r in out_rows if r.get("shop")})
        return {
            "ok": True,
            "sheet": name,
            "header_row": header_row,
            "row_count": len(out_rows),
            "fba_count": len(by_fba),
            "sku_count": len({x["sku"] for x in out_rows}),
            "shops": shops,
            "shipment_weights": shipment_weights,
            "rows": out_rows,
            "source": "mapped",
        }
    finally:
        wb.close()
